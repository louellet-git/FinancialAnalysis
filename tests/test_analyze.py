import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from datetime import date

import pytest

from financial_analysis.schema import CompanyFinancials, PeriodFinancials
from scripts.analyze import _flatten, write_xlsx, COLUMNS


def _make_company() -> CompanyFinancials:
    p1 = PeriodFinancials(
        ticker="TEST",
        period_end=date(2024, 1, 28),
        form="10-K",
        filed_date=date(2024, 2, 21),
        revenue=60_920_000_000,
        operating_cash_flow=28_090_000_000,
        capex=-1_070_000_000,
        total_assets=65_730_000_000,
        raw_extras={"us-gaap/PaymentsToAcquireEquitySecuritiesFvNi": 860_000_000},
    )
    p2 = PeriodFinancials(
        ticker="TEST",
        period_end=date(2023, 10, 29),
        form="10-Q",
        filed_date=date(2023, 11, 21),
        revenue=18_120_000_000,
        raw_extras={},
    )
    return CompanyFinancials(ticker="TEST", cik="1045810", company_name="Test Co", periods=[p1, p2])


def test_flatten_returns_all_tabs():
    company = _make_company()
    tabs = _flatten(company)
    assert set(tabs.keys()) == {
        "Income Statement", "Balance Sheet", "Cash Flow", "Raw",
        "Prices", "Estimates", "Estimates Fwd", "Options", "Options Summary",
    }


def test_flatten_income_statement_has_revenue_row():
    tabs = _flatten(_make_company())
    is_rows = tabs["Income Statement"]
    metrics = [r["metric"] for r in is_rows]
    assert "revenue" in metrics


def test_flatten_cash_flow_has_free_cash_flow():
    tabs = _flatten(_make_company())
    cf_rows = tabs["Cash Flow"]
    metrics = [r["metric"] for r in cf_rows]
    assert "free_cash_flow" in metrics


def test_flatten_raw_tab_has_unmapped_concepts():
    tabs = _flatten(_make_company())
    raw_rows = tabs["Raw"]
    metrics = [r["metric"] for r in raw_rows]
    assert "us-gaap/PaymentsToAcquireEquitySecuritiesFvNi" in metrics


def test_flatten_none_values_not_included():
    tabs = _flatten(_make_company())
    for tab_name, tab_rows in tabs.items():
        for row in tab_rows:
            assert row.get("value") is not None, f"None value in {tab_name}: {row['metric']}"


def test_flatten_row_has_all_columns():
    tabs = _flatten(_make_company())
    for tab_name, tab_rows in tabs.items():
        for row in tab_rows:
            for col in COLUMNS:
                assert col in row, f"Column '{col}' missing from row in {tab_name}"


def test_flatten_capex_unit_is_usd():
    tabs = _flatten(_make_company())
    capex_rows = [r for r in tabs["Cash Flow"] if r["metric"] == "capex"]
    assert all(r["unit"] == "USD" for r in capex_rows)


def test_write_xlsx_creates_file_with_all_tabs(tmp_path):
    import openpyxl
    company = _make_company()
    out = tmp_path / "TEST_financials.xlsx"
    write_xlsx(company, out)
    assert out.exists()
    wb = openpyxl.load_workbook(out)
    assert set(wb.sheetnames) == {
        "Income Statement", "Balance Sheet", "Cash Flow", "Raw",
        "Prices", "Estimates", "Estimates Fwd", "Options", "Options Summary",
    }


def test_write_xlsx_has_header_row(tmp_path):
    import openpyxl
    company = _make_company()
    out = tmp_path / "TEST_financials.xlsx"
    write_xlsx(company, out)
    wb = openpyxl.load_workbook(out)
    ws = wb["Income Statement"]
    header = [cell.value for cell in ws[1]]
    assert header == COLUMNS


def test_write_xlsx_returns_total_row_count(tmp_path):
    company = _make_company()
    out = tmp_path / "TEST_financials.xlsx"
    total = write_xlsx(company, out)
    assert total > 0


def _attach_option_chain(company: CompanyFinancials) -> CompanyFinancials:
    from financial_analysis.schema import OptionChainSnapshot, OptionQuote
    company.option_chain = OptionChainSnapshot(
        ticker="TEST",
        snapshot_date=date(2026, 7, 16),
        spot_price=100.0,
        quotes=[
            OptionQuote("TEST260724C00100000", "call", date(2026, 7, 24), 100.0,
                        last_price=5.0, bid=4.9, ask=5.1, volume=10, open_interest=100,
                        implied_volatility=0.30, in_the_money=False),
            OptionQuote("TEST260724P00100000", "put", date(2026, 7, 24), 100.0,
                        last_price=4.8, bid=4.7, ask=4.9, volume=20, open_interest=50,
                        implied_volatility=0.32, in_the_money=False),
        ],
    )
    return company


def test_flatten_options_tab_one_row_per_contract():
    tabs = _flatten(_attach_option_chain(_make_company()))
    assert len(tabs["Options"]) == 2
    types = {r["option_type"] for r in tabs["Options"]}
    assert types == {"call", "put"}
    assert tabs["Options"][0]["expiry"] == "2026-07-24"


def test_flatten_options_summary_has_per_expiry_and_total_rows():
    tabs = _flatten(_attach_option_chain(_make_company()))
    rows = tabs["Options Summary"]
    expiries = [r["expiry"] for r in rows]
    assert expiries == ["2026-07-24", "ALL"]
    per_exp = rows[0]
    assert per_exp["atm_strike"] == 100.0
    assert per_exp["atm_iv"] == pytest.approx(0.31)
    assert per_exp["implied_move"] == pytest.approx(0.098)
    assert rows[1]["put_call_volume_ratio"] == pytest.approx(2.0)


def test_flatten_no_option_chain_gives_empty_tabs():
    tabs = _flatten(_make_company())
    assert tabs["Options"] == []
    assert tabs["Options Summary"] == []


def _attach_consensus(company: CompanyFinancials) -> CompanyFinancials:
    from financial_analysis.schema import ConsensusEstimates
    company.consensus = ConsensusEstimates(
        ticker="TEST",
        snapshot_date=date(2026, 7, 17),
        revenue_estimate={"0y": {"avg": 66.5e9, "numberOfAnalysts": 18.0, "growth": None}},
        earnings_estimate={"0y": {"avg": 10.95}},
        analyst_price_targets={"mean": 271.3, "high": None},
    )
    return company


def test_flatten_estimates_fwd_rows_from_consensus():
    tabs = _flatten(_attach_consensus(_make_company()))
    rows = tabs["Estimates Fwd"]
    kinds = {r["kind"] for r in rows}
    assert kinds == {"revenue_estimate", "earnings_estimate", "price_target"}
    rev_avg = next(r for r in rows if r["kind"] == "revenue_estimate" and r["metric"] == "avg")
    assert rev_avg["period"] == "0y"
    assert rev_avg["value"] == 66.5e9
    assert rev_avg["snapshot_date"] == "2026-07-17"
    # None values (no growth figure, no high target) are dropped, not written
    assert all(r["value"] is not None for r in rows)
    assert len(rows) == 4  # avg + numberOfAnalysts (revenue), avg (eps), mean (target)


def test_flatten_no_consensus_gives_empty_estimates_fwd_tab():
    tabs = _flatten(_make_company())
    assert tabs["Estimates Fwd"] == []


def test_flatten_rpo_lands_on_balance_sheet_tab_as_disclosure():
    company = _make_company()
    company.periods[0].remaining_performance_obligation = 71_000_000_000
    tabs = _flatten(company)
    rpo_rows = [r for r in tabs["Balance Sheet"] if r["metric"] == "remaining_performance_obligation"]
    assert len(rpo_rows) == 1
    assert rpo_rows[0]["statement"] == "disclosure"
    assert rpo_rows[0]["value"] == 71_000_000_000
